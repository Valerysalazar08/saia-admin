
import os, threading
from datetime import datetime
from pathlib import Path
from PyQt6.QtWidgets import QWidget, QFrame, QLabel, QPushButton, QHBoxLayout, QVBoxLayout
from PyQt6.QtCore import Qt, QObject, QTimer, pyqtSignal

from app.ui.theme import (
    BG_APP, BG_CARD, BG_INPUT, BORDER, PRIMARY, ERROR, ERROR_BG,
    SUCCESS_BG, SUCCESS_TEXT, WARNING, WARNING_BG,
    TEXT_MUTED, TEXT_PRIMARY, INFO, INFO_BG,
    CARD_RADIUS, BADGE_RADIUS, font,
)
from app.ui.atoms.inputs   import SearchInput, TextInput, Dropdown
from app.ui.atoms.labels   import Heading
from app.ui.molecules.data_table import DataTable
from app.models.auditoria_model     import AuditoriaModel
from app.config.settings            import APP


_ACCION_COLORS = {
    "CREAR":      (SUCCESS_BG, SUCCESS_TEXT),
    "ACTUALIZAR": (INFO_BG,    INFO),
    "BLOQUEAR":   (ERROR_BG,   ERROR),
    "HABILITAR":  (SUCCESS_BG, SUCCESS_TEXT),
    "ELIMINAR":   (ERROR_BG,   ERROR),
    "LOGIN":      (INFO_BG,    PRIMARY),
    "LOGOUT":     (WARNING_BG, WARNING),
}


class _Sig(QObject):
    done = pyqtSignal(list)


class AuditoriaView(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)

        self._last_data = []
        self._updating_filters = False

        self.setStyleSheet(f"background:{BG_APP};")
        self._build()

        self._load_filter_options()
        self.load_data()

        self._timer = QTimer(self)
        self._timer.timeout.connect(self.load_data)
        self._timer.start(30000)

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 20, 24, 8); lay.setSpacing(0)

        # Toolbar
        tb = QWidget(); tb.setStyleSheet("background:transparent;")
        tbl = QHBoxLayout(tb); tbl.setContentsMargins(0,0,0,0); tbl.setSpacing(8)
        tbl.addWidget(Heading("Historial de Auditoría", level=2), stretch=1)


        refresh_btn = QPushButton_("↻  Actualizar", PRIMARY, "#F0FAFD", self.load_data)
        pdf_btn = QPushButton_("PDF", ERROR, "#FFF1F2", lambda: self._export("pdf"))
        xls_btn = QPushButton_("Excel", "#168A62", "#EDFCF6", lambda: self._export("excel"))
        tbl.addWidget(refresh_btn)
        tbl.addWidget(pdf_btn); tbl.addWidget(xls_btn)
        lay.addWidget(tb); lay.addSpacing(10)


        fc = QFrame()

        fc.setObjectName("AuditFilterCard")
        fc.setStyleSheet(f"""
            QFrame#AuditFilterCard {{
                background: {BG_CARD};
                border-radius: {CARD_RADIUS}px;
                border: 1px solid {BORDER};
            }}
        """)
        fl = QHBoxLayout(fc); fl.setContentsMargins(16, 12, 16, 12); fl.setSpacing(12)

        self._search = SearchInput("Buscar en descripción...", width=220,
                                   on_change=self._on_search)
        fl.addLayout(self._filter_field("Buscar registro", self._search), stretch=2)

        self._accion_dd = Dropdown(
            ["Todas"],
            width=120,
            command=lambda _: self._filter_changed()
        )
        fl.addLayout(self._filter_field("Acción", self._accion_dd))

        self._entidad_dd = Dropdown(
            ["Todas"],
            width=120,
            command=lambda _: self._filter_changed()
        )
        fl.addLayout(self._filter_field("Entidad", self._entidad_dd))

        self._fecha_ini = TextInput("YYYY-MM-DD", width=110, height=36)
        fl.addLayout(self._filter_field("Desde", self._fecha_ini))
        self._fecha_fin = TextInput("YYYY-MM-DD", width=110, height=36)
        fl.addLayout(self._filter_field("Hasta", self._fecha_fin))

        fb = QPushButton_("Filtrar", PRIMARY, "#F0FAFD", self.load_data)
        fl.addWidget(fb, alignment=Qt.AlignmentFlag.AlignBottom)
        lay.addWidget(fc); lay.addSpacing(8)

        cols = [
            {"key":"fecha_hora",       "header":"Fecha/Hora",    "width":170},
            {"key":"tipo_accion",      "header":"Acción",        "width":120,
             "renderer":self._render_accion},
            {"key":"entidad",          "header":"Entidad",       "width":110},
            {"key":"descripcion",      "header":"Descripción",   "width":390,
             "header_width":422, "wrap":True},
            {"key":"afectado_nombres", "header":"Afectado",      "width":190},
            {"key":"usuario",          "header":"Realizado por", "width":220},
        ]
        self._table = DataTable(columns=cols, row_height=66, page_size=12)
        lay.addWidget(self._table, stretch=1); lay.addSpacing(4)

        self._status = QLabel("")
        self._status.setFont(font(11))
        self._status.setStyleSheet(f"color:{TEXT_MUTED}; background:transparent;")
        lay.addWidget(self._status)

    def _lbl(self, t):
        l = QLabel(t); l.setFont(font(11))
        l.setStyleSheet(f"color:{TEXT_MUTED}; background:transparent;"); return l

    def _filter_field(self, label: str, widget: QWidget) -> QVBoxLayout:
        """Etiqueta"""
        lay = QVBoxLayout()
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(4)
        caption = QLabel(label)
        caption.setFont(font(9, bold=True))
        caption.setStyleSheet(
            f"color:{TEXT_MUTED}; background:transparent; border:none; padding:0;")
        lay.addWidget(caption)
        lay.addWidget(widget)
        return lay

    def _render_accion(self, parent, row, val):
        txt = str(val).upper()
        bg, fg = _ACCION_COLORS.get(txt, (BG_INPUT, TEXT_MUTED))
        w = QWidget(parent); w.setStyleSheet("background:transparent;")
        h = QHBoxLayout(w); h.setContentsMargins(4,0,0,0)
        lbl = QLabel(txt); lbl.setFont(font(10, bold=True))
        lbl.setStyleSheet(f"""
            color:{fg}; background:{bg}; border-radius:{BADGE_RADIUS}px;
            padding:2px 8px;
        """)
        lbl.setFixedHeight(22)
        h.addWidget(lbl); h.addStretch(); return w

    def _filter_changed(self):
        if not self._updating_filters:
            self.load_data()


    def _load_filter_options(self):
        try:
            self._updating_filters = True

            acciones = AuditoriaModel.get_acciones()
            entidades = AuditoriaModel.get_entidades()

            self._accion_dd.set_options(acciones)
            self._entidad_dd.set_options(entidades)

        except Exception:
            pass

        finally:
            self._updating_filters = False

    def load_data(self, search: str = None):
        s  = search if search is not None else self._search.get()
        ac = self._accion_dd.get(); en = self._entidad_dd.get()
        fi = self._fecha_ini.get().strip() or None
        ff = self._fecha_fin.get().strip() or None
        sig = _Sig(self); sig.done.connect(self._on_data)
        def fetch():
            try:
                data = AuditoriaModel.get_all(search=s, accion=ac, entidad=en,
                                               fecha_inicio=fi, fecha_fin=ff)
                for row in data:
                    n = row.get("afectado_nombres") or ""
                    a = row.get("afectado_ape") or ""
                    row["afectado_nombres"] = f"{n} {a}".strip() or str(row.get("num_doc",""))
            except Exception: data = []
            sig.done.emit(data)
        threading.Thread(target=fetch, daemon=True).start()

    def _on_data(self, data):
        self._last_data = data
        self._table.load(data)
        self._status.setText(f"{len(data)} registro(s)")

    def _on_search(self, text): self.load_data(text)

    def _export(self, fmt: str):
        if not self._last_data:
            self._status.setText("No hay datos para exportar."); return
        headers    = ["fecha_hora","tipo_accion","entidad","descripcion","afectado_nombres","usuario"]
        col_titles = ["Fecha/Hora","Acción","Entidad","Descripción","Afectado","Realizado por"]
        ts      = datetime.now().strftime("%Y%m%d_%H%M%S")
        exports = Path(APP["exports_dir"]); exports.mkdir(exist_ok=True)
        try:
            if fmt == "pdf":
                path = exports / f"auditoria_{ts}.pdf"
                self._export_pdf(self._last_data, headers, col_titles, str(path))
            else:
                path = exports / f"auditoria_{ts}.xlsx"
                self._export_excel(self._last_data, headers, col_titles, str(path))
            self._status.setText(f"✓ Exportado: {path.name}")
            os.startfile(str(path))
        except Exception as e:
            self._status.setText(f"Error al exportar: {e}")

    def _export_pdf(self, data, headers, col_titles, path):
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                                leftMargin=1.5*cm, rightMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        story  = []
        ts = ParagraphStyle("t", parent=styles["Title"],
                            textColor=colors.HexColor("#33BEDC"), fontSize=14)
        story.append(Paragraph("SAIA — Historial de Auditoría", ts))
        story.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Registros: {len(data)}",
                               styles["Normal"]))
        story.append(Spacer(1, 0.4*cm))
        pw = landscape(A4)[0] - 3*cm
        cw = [pw*r for r in [0.13,0.09,0.08,0.35,0.18,0.17]]
        td = [col_titles] + [[str(r.get(h,"") or "") for h in headers] for r in data]
        t = Table(td, colWidths=cw, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#33BEDC")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),8),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.HexColor("#F9FAFB"),colors.HexColor("#FFFFFF")]),
            ("TEXTCOLOR",(0,1),(-1,-1),colors.HexColor("#1A1A2E")),("FONTSIZE",(0,1),(-1,-1),7),
            ("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#E5E7EB")),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("TOPPADDING",(0,0),(-1,-1),4),
            ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ]))
        story.append(t); doc.build(story)

    def _export_excel(self, data, headers, col_titles, path):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Auditoría"
        hf = PatternFill("solid", fgColor="33BEDC"); hfnt = Font(color="FFFFFF", bold=True, size=10)
        ef = PatternFill("solid", fgColor="F9FAFB"); of = PatternFill("solid", fgColor="FFFFFF")
        cf = Font(color="1A1A2E", size=9); thin = Side(style="thin", color="E5E7EB")
        brd = Border(left=thin, right=thin, top=thin, bottom=thin)
        for ci, t in enumerate(col_titles, 1):
            cell = ws.cell(1,ci,t); cell.font=hfnt; cell.fill=hf
            cell.alignment=Alignment(horizontal="center",vertical="center"); cell.border=brd
        for ri, row in enumerate(data, 2):
            fill = ef if ri%2==0 else of
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(ri,ci,str(row.get(h,"") or ""))
                cell.font=cf; cell.fill=fill
                cell.alignment=Alignment(vertical="center",wrap_text=True); cell.border=brd
        for col in ws.columns:
            ml = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(ml+3, 45)
        ws.freeze_panes="A2"; wb.save(path)

    def closeEvent(self, event):
        self._timer.stop(); super().closeEvent(event)


def QPushButton_(text, color, bg, fn):
    """Botón de acción ligero para la barra superior de Auditoría."""
    b = QPushButton(text); b.setFixedHeight(34); b.setMinimumWidth(78)
    b.setFont(font(11, bold=True)); b.setCursor(Qt.CursorShape.PointingHandCursor)
    b.setFocusPolicy(Qt.FocusPolicy.NoFocus)
    b.setStyleSheet(f"""
        QPushButton {{background:{bg}; color:{color};
                      border:1px solid {BORDER}; border-radius:7px; padding:0 13px; outline:none;}}
        QPushButton:hover {{background:{BG_INPUT}; border-color:{color};}}
        QPushButton:pressed {{background:{bg}; border-color:{color};}}
        QPushButton:focus {{outline:none;}}
    """)
    b.clicked.connect(fn); return b
