"""
VISTA Qt — Generador de reportes.
"""
import os
from datetime import datetime
from pathlib import Path
from PyQt6.QtWidgets import (
    QWidget, QFrame, QLabel, QTextEdit, QPushButton,
    QVBoxLayout, QHBoxLayout, QScrollArea,
)
from PyQt6.QtCore import Qt

from app.ui.theme import (
    BG_APP, BG_CARD, BG_INPUT, BG_HOVER, BORDER,
    TEXT_PRIMARY, TEXT_SECONDARY, TEXT_MUTED, PRIMARY,
    CARD_RADIUS, font,
)
from app.ui.atoms.buttons import PrimaryButton, SecondaryButton
from app.ui.atoms.inputs  import TextInput, Dropdown
from app.ui.atoms.labels  import Heading, Divider
from app.config.settings     import APP


class ReportesView(QWidget):
    TIPOS = ["Historial de ingresos","Lista de aprendices",
             "Lista de guardas","Ingresos por portería"]

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setStyleSheet(f"background:{BG_APP};")
        self._build()
        self._refresh_files()

    def _build(self):
        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 20, 24, 16); lay.setSpacing(0)

        lay.addWidget(Heading("Generador de Reportes", level=2))
        sub = QLabel("Genera y exporta reportes en PDF, Excel o CSV.")
        sub.setFont(font(11)); sub.setStyleSheet(f"color:{TEXT_MUTED}; background:transparent;")
        sub.setContentsMargins(0,4,0,16); lay.addWidget(sub)

        # ── Card configuración ────────────────────────────────────────────────
        card = QFrame()
        card.setStyleSheet(f"QFrame{{background:{BG_CARD};border-radius:{CARD_RADIUS}px;border:1px solid {BORDER};}}")
        cl = QVBoxLayout(card); cl.setContentsMargins(16,14,16,14); cl.setSpacing(8)

        tl = QLabel("Configurar reporte"); tl.setFont(font(14, bold=True))
        tl.setStyleSheet(f"color:{TEXT_PRIMARY}; background:transparent;")
        cl.addWidget(tl); cl.addWidget(Divider())

        row = QWidget(); row.setStyleSheet("background:transparent;")
        rl  = QHBoxLayout(row); rl.setContentsMargins(0,8,0,0); rl.setSpacing(12)

        for label_text, attr, values in [
            ("Tipo de reporte",  "_tipo_dd",    self.TIPOS),
            ("Formato",          "_formato_dd", ["PDF","Excel","CSV"]),
        ]:
            col = QWidget(); col.setStyleSheet("background:transparent;")
            vl  = QVBoxLayout(col); vl.setContentsMargins(0,0,0,0); vl.setSpacing(4)
            lbl = QLabel(label_text); lbl.setFont(font(11, bold=True))
            lbl.setStyleSheet(f"color:{TEXT_SECONDARY}; background:transparent;")
            vl.addWidget(lbl)
            dd = Dropdown(values, width=200)
            setattr(self, attr, dd); vl.addWidget(dd); rl.addWidget(col)

        for label_text, attr in [("Fecha inicio","_fecha_ini"),("Fecha fin","_fecha_fin")]:
            col = QWidget(); col.setStyleSheet("background:transparent;")
            vl  = QVBoxLayout(col); vl.setContentsMargins(0,0,0,0); vl.setSpacing(4)
            lbl = QLabel(label_text); lbl.setFont(font(11, bold=True))
            lbl.setStyleSheet(f"color:{TEXT_SECONDARY}; background:transparent;")
            vl.addWidget(lbl)
            ti = TextInput("YYYY-MM-DD", width=130, height=42)
            setattr(self, attr, ti); vl.addWidget(ti); rl.addWidget(col)

        rl.addStretch(); cl.addWidget(row)

        btn_row = QWidget(); btn_row.setStyleSheet("background:transparent;")
        brl = QHBoxLayout(btn_row); brl.setContentsMargins(0,4,0,0); brl.setSpacing(8)
        gen = PrimaryButton("📄 Generar reporte", width=180, height=38)
        gen.clicked.connect(self._generate)
        folder = SecondaryButton("📂 Abrir carpeta", width=150, height=38)
        folder.clicked.connect(lambda: os.startfile(APP["exports_dir"]))
        brl.addWidget(gen); brl.addWidget(folder); brl.addStretch()
        cl.addWidget(btn_row)
        lay.addWidget(card); lay.addSpacing(16)

        # ── Log ───────────────────────────────────────────────────────────────
        self._log = QTextEdit()
        self._log.setReadOnly(True)
        self._log.setFixedHeight(130)
        self._log.setFont(font(10, family="Consolas"))
        self._log.setStyleSheet(f"""
            QTextEdit {{background:{BG_CARD}; border:1px solid {BORDER};
                        border-radius:{CARD_RADIUS}px; color:{TEXT_SECONDARY};}}
        """)
        lay.addWidget(self._log)
        lay.addSpacing(12)

        # ── Lista de archivos ─────────────────────────────────────────────────
        lay.addWidget(QLabel("Reportes generados") if False else
                      self._section_label("Reportes generados"))

        self._files_scroll = QScrollArea()
        self._files_scroll.setWidgetResizable(True)
        self._files_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self._files_scroll.setFixedHeight(180)
        self._files_scroll.setStyleSheet(f"background:{BG_CARD}; border-radius:{CARD_RADIUS}px;")

        self._files_body = QWidget()
        self._files_body.setStyleSheet(f"background:{BG_CARD};")
        self._files_lay = QVBoxLayout(self._files_body)
        self._files_lay.setContentsMargins(8,8,8,8); self._files_lay.setSpacing(3)
        self._files_lay.addStretch()
        self._files_scroll.setWidget(self._files_body)
        lay.addWidget(self._files_scroll, stretch=1)

    def _section_label(self, text):
        l = QLabel(text); l.setFont(font(14, bold=True))
        l.setStyleSheet(f"color:{TEXT_PRIMARY}; background:transparent;")
        l.setContentsMargins(0,0,0,4); return l

    def _log_msg(self, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        self._log.append(f"[{ts}] {msg}")

    def _generate(self):
        fmt  = self._formato_dd.get()
        tipo = self._tipo_dd.get()
        self._log_msg(f"Generando '{tipo}' en {fmt}…")
        try:
            data, headers = self._get_data()
        except Exception as e:
            self._log_msg(f"Error: {e}"); return
        if not data:
            self._log_msg("No hay datos."); return

        ts      = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe    = tipo.replace(" ","_").lower()
        exports = Path(APP["exports_dir"]); exports.mkdir(exist_ok=True)
        try:
            if fmt == "PDF":
                path = exports / f"{safe}_{ts}.pdf"
                self._export_pdf(data, headers, str(path), tipo)
            elif fmt == "Excel":
                path = exports / f"{safe}_{ts}.xlsx"
                self._export_excel(data, headers, str(path))
            else:
                path = exports / f"{safe}_{ts}.csv"
                self._export_csv(data, headers, str(path))
            self._log_msg(f"✓ Guardado: {path.name}")
            self._refresh_files()
        except Exception as e:
            self._log_msg(f"Error al generar: {e}")

    def _get_data(self):
        from app.models.historial_model  import HistorialModel
        from app.models.aprendiz_model   import AprendizSaiaModel
        from app.models.guarda_model     import GuardaModel
        tipo = self._tipo_dd.get()
        fi   = self._fecha_ini.get().strip() or None
        ff   = self._fecha_fin.get().strip() or None
        if tipo == "Historial de ingresos":
            return HistorialModel.get_all(fecha_inicio=fi, fecha_fin=ff, limit=5000), \
                   ["id_ingreso","fecha_hora_ingreso","fecha_hora_salida","nombres","p_ape",
                    "num_doc","porteria","estado_movimiento","guarda_nombres"]
        elif tipo == "Lista de aprendices":
            return AprendizSaiaModel.get_all(), \
                   ["num_doc","tip_doc","nombres","p_ape","email","tel","cuenta_estado"]
        elif tipo == "Lista de guardas":
            return GuardaModel.get_all(), \
                   ["id_guarda","num_doc","nombres","p_ape","turno","empresa_seg","tel"]
        elif tipo == "Ingresos por portería":
            return HistorialModel.ingresos_por_porteria(), ["porteria","total"]
        return [], []

    def _export_pdf(self, data, headers, path, title):
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        doc = SimpleDocTemplate(path, pagesize=landscape(A4),
                                leftMargin=1.5*cm, rightMargin=1.5*cm,
                                topMargin=2*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet(); story = []
        ts_s = ParagraphStyle("t", parent=styles["Title"],
                              textColor=colors.HexColor("#33BEDC"), fontSize=16)
        story.append(Paragraph(f"SAIA — {title}", ts_s))
        story.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                               styles["Normal"]))
        story.append(Spacer(1, 0.5*cm))
        ch = [h.replace("_"," ").title() for h in headers]
        td = [ch] + [[str(r.get(h,"") or "") for h in headers] for r in data]
        cw = (landscape(A4)[0]-3*cm)/len(headers)
        t = Table(td, colWidths=[cw]*len(headers), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#33BEDC")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
            ("FONTSIZE",(0,0),(-1,0),8),("FONTSIZE",(0,1),(-1,-1),7),
            ("GRID",(0,0),(-1,-1),0.3,colors.HexColor("#E5E7EB")),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ]))
        story.append(t); doc.build(story)

    def _export_excel(self, data, headers, path):
        import openpyxl; from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook(); ws = wb.active; ws.title="Reporte"
        hf = PatternFill("solid",fgColor="33BEDC"); hfnt = Font(color="FFFFFF",bold=True,size=10)
        cf = Font(size=9); thin = Side(style="thin",color="E5E7EB")
        brd = Border(left=thin,right=thin,top=thin,bottom=thin)
        for ci, h in enumerate([h.replace("_"," ").title() for h in headers], 1):
            c=ws.cell(1,ci,h); c.font=hfnt; c.fill=hf; c.border=brd
            c.alignment=Alignment(horizontal="center",vertical="center")
        for ri, row in enumerate(data, 2):
            for ci, h in enumerate(headers, 1):
                c=ws.cell(ri,ci,str(row.get(h,"") or "")); c.font=cf; c.border=brd
        ws.freeze_panes="A2"; wb.save(path)

    def _export_csv(self, data, headers, path):
        import csv
        with open(path,"w",newline="",encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
            w.writeheader(); w.writerows(data)

    def _refresh_files(self):
        lay = self._files_lay
        while lay.count() > 1:
            item = lay.takeAt(0)
            if item.widget(): item.widget().deleteLater()

        exports = Path(APP["exports_dir"])
        files = sorted(exports.glob("*"), key=lambda p: p.stat().st_mtime, reverse=True)
        if not files:
            lbl = QLabel("Sin reportes generados aún."); lbl.setFont(font(11))
            lbl.setStyleSheet(f"color:{TEXT_MUTED}; background:transparent;")
            lay.insertWidget(0, lbl); return

        ext_icons = {".pdf":"📄",".xlsx":"📊",".csv":"📋"}
        for f in files[:20]:
            item = QFrame()
            item.setStyleSheet(f"QFrame{{background:{BG_INPUT};border-radius:6px;border:none;}}")
            rh = QHBoxLayout(item); rh.setContentsMargins(8,4,8,4); rh.setSpacing(8)
            ic = QLabel(ext_icons.get(f.suffix.lower(),"📁"))
            ic.setFont(font(14)); ic.setStyleSheet("background:transparent;"); rh.addWidget(ic)
            nl = QLabel(f.name); nl.setFont(font(11))
            nl.setStyleSheet(f"color:{TEXT_PRIMARY}; background:transparent;")
            rh.addWidget(nl, stretch=1)
            sl = QLabel(f"{f.stat().st_size/1024:.1f} KB"); sl.setFont(font(10))
            sl.setStyleSheet(f"color:{TEXT_MUTED}; background:transparent;"); rh.addWidget(sl)
            ob = QPushButton("📂"); ob.setFixedSize(26,26)
            ob.setStyleSheet(f"background:transparent;border:none;font-size:14px;color:{PRIMARY};")
            ob.setCursor(Qt.CursorShape.PointingHandCursor)
            ob.clicked.connect(lambda _, p=str(f.parent): os.startfile(p))
            rh.addWidget(ob)
            lay.insertWidget(lay.count()-1, item)
